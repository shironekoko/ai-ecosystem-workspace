"""
Convert Markdown Report to Microsoft Word (.docx) Document

วิธีใช้งาน:
    uv run python utils/convert_report_to_word.py

ผลลัพธ์:
    project_architecture_report.docx (ในโฟลเดอร์ ai-ecosystem-workspace)
"""

import os
import re
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

if sys.platform == "win32":
    sys.stdout.reconfigure(encoding="utf-8")


def set_cell_background(cell, fill_hex):
    """กำหนดสีพื้นหลังของ Cell ในตาราง"""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill_hex)
    tc_pr.append(shd)


def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """กำหนดระยะขอบใน Cell"""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for margin, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{margin}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)


def convert_md_to_docx(md_filepath, docx_filepath):
    doc = docx.Document()

    # ตั้งค่า Page Margins (1 นิ้วทุกด้าน)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    # อ่านเนื้อหาจาก Markdown
    with open(md_filepath, 'r', encoding='utf-8') as f:
        content = f.read()

    lines = content.split('\n')
    i = 0
    in_code_block = False
    code_lines = []
    code_lang = ""

    while i < len(lines):
        line = lines[i]

        # ── Code Block ──
        if line.strip().startswith('```'):
            if not in_code_block:
                in_code_block = True
                code_lang = line.strip().replace('```', '')
                code_lines = []
            else:
                in_code_block = False
                # เขียน Code Block เป็นตารางเซลล์เดียวพื้นหลังสีเทาอ่อน
                table = doc.add_table(rows=1, cols=1)
                table.alignment = WD_TABLE_ALIGNMENT.CENTER
                cell = table.cell(0, 0)
                set_cell_background(cell, 'F4F6F8')
                set_cell_margins(cell, top=140, bottom=140, left=200, right=200)

                p = cell.paragraphs[0]
                p.paragraph_format.space_before = Pt(4)
                p.paragraph_format.space_after = Pt(4)
                p.paragraph_format.line_spacing = 1.15
                
                code_text = '\n'.join(code_lines)
                run = p.add_run(code_text)
                run.font.name = 'Consolas'
                run.font.size = Pt(9.5)
                run.font.color.rgb = RGBColor(0x24, 0x29, 0x2E)

                # Add empty paragraph after table
                p_space = doc.add_paragraph()
                p_space.paragraph_format.space_before = Pt(0)
                p_space.paragraph_format.space_after = Pt(4)

            i += 1
            continue

        if in_code_block:
            code_lines.append(line)
            i += 1
            continue

        # ── Markdown Table ──
        if '|' in line and i + 1 < len(lines) and '|' in lines[i + 1] and '---' in lines[i + 1]:
            table_lines = []
            while i < len(lines) and '|' in lines[i]:
                # ข้ามบรรทัดแยกหัวตาราง (| --- | --- |)
                if '---' in lines[i]:
                    i += 1
                    continue
                table_lines.append(lines[i])
                i += 1

            if table_lines:
                # Parse Table Rows
                parsed_rows = []
                for tline in table_lines:
                    cols = [c.strip() for c in tline.split('|')[1:-1]]
                    if cols:
                        parsed_rows.append(cols)

                if parsed_rows:
                    num_cols = max(len(row) for row in parsed_rows)
                    tbl = doc.add_table(rows=len(parsed_rows), cols=num_cols)
                    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                    tbl.style = 'Table Grid'

                    for r_idx, row in enumerate(parsed_rows):
                        for c_idx, cell_value in enumerate(row):
                            if c_idx < num_cols:
                                cell = tbl.cell(r_idx, c_idx)
                                cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
                                set_cell_margins(cell, top=100, bottom=100, left=150, right=150)
                                p = cell.paragraphs[0]
                                p.paragraph_format.space_before = Pt(2)
                                p.paragraph_format.space_after = Pt(2)
                                p.paragraph_format.line_spacing = 1.15

                                # Clean markdown bold/code from cell text
                                clean_val = re.sub(r'<br\s*/?>', '\n', cell_value)
                                clean_val = clean_val.replace('**', '').replace('`', '')
                                
                                run = p.add_run(clean_val)
                                run.font.name = 'Cordia New'
                                run.font.size = Pt(13)

                                # Header Row Styling
                                if r_idx == 0:
                                    set_cell_background(cell, '1F4E79') # Dark Blue Fill
                                    run.font.bold = True
                                    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                                else:
                                    if r_idx % 2 == 1:
                                        set_cell_background(cell, 'F9FAFB')
                                    else:
                                        set_cell_background(cell, 'FFFFFF')
                                    run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)

                    # Spacing after table
                    sp = doc.add_paragraph()
                    sp.paragraph_format.space_before = Pt(0)
                    sp.paragraph_format.space_after = Pt(6)

            continue

        # ── Headings ──
        if line.startswith('# '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(16)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run(line[2:].strip())
            run.font.name = 'Angsana New'
            run.font.size = Pt(24)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1F, 0x4E, 0x79)
            i += 1
            continue

        if line.startswith('## '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(line[3:].strip())
            run.font.name = 'Angsana New'
            run.font.size = Pt(18)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x2E, 0x75, 0xB6)
            i += 1
            continue

        if line.startswith('### '):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(line[4:].strip())
            run.font.name = 'Angsana New'
            run.font.size = Pt(16)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1B, 0x36, 0x5D)
            i += 1
            continue

        # ── Bullet Lists ──
        if line.strip().startswith('- ') or line.strip().startswith('* '):
            p = doc.add_paragraph(style='List Bullet')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            raw_text = re.sub(r'^[\-\*]\s+', '', line.strip())
            _add_formatted_runs(p, raw_text)
            i += 1
            continue

        # ── Numbered Lists ──
        num_match = re.match(r'^(\d+)\.\s+(.*)$', line.strip())
        if num_match:
            p = doc.add_paragraph(style='List Number')
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(2)
            p.paragraph_format.line_spacing = 1.15
            
            _add_formatted_runs(p, num_match.group(2))
            i += 1
            continue

        # ── Horizontal Rule ──
        if line.strip() == '---':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run('_________________________________________________________________________________')
            run.font.color.rgb = RGBColor(0xD3, 0xD3, 0xD3)
            i += 1
            continue

        # ── Normal Paragraph ──
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)
            p.paragraph_format.line_spacing = 1.15
            _add_formatted_runs(p, line.strip())
        
        i += 1

    doc.save(docx_filepath)
    print(f"✅ บันทึกเอกสาร Word สำเร็จที่: {os.path.abspath(docx_filepath)}")


def _add_formatted_runs(paragraph, text):
    """ช่วยแกะ Markdown Bold (`**text**`) และ Code (`code`) เพื่อใส่ style ใน Word"""
    # Regex หา **bold** และ `code`
    tokens = re.split(r'(\*\*.*?\*\*|`.*?`)', text)
    for token in tokens:
        if not token:
            continue
        if token.startswith('**') and token.endswith('**'):
            run = paragraph.add_run(token[2:-2])
            run.font.name = 'Cordia New'
            run.font.size = Pt(14)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)
        elif token.startswith('`') and token.endswith('`'):
            run = paragraph.add_run(token[1:-1])
            run.font.name = 'Consolas'
            run.font.size = Pt(11)
            run.font.color.rgb = RGBColor(0xC7, 0x25, 0x4E) # Reddish code color
        else:
            run = paragraph.add_run(token)
            run.font.name = 'Cordia New'
            run.font.size = Pt(14)
            run.font.color.rgb = RGBColor(0x26, 0x26, 0x26)


if __name__ == "__main__":
    src_md = os.path.join(os.path.dirname(__file__), '..', '..', 'project_architecture_report.md')
    dst_docx = os.path.join(os.path.dirname(__file__), '..', '..', 'project_architecture_report.docx')
    convert_md_to_docx(src_md, dst_docx)
